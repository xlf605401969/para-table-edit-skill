#!/usr/bin/env python3
"""Edit and validate ParameterTable.xlsx workbooks."""

from __future__ import annotations

import argparse
import json
import re
import sys
import warnings
from copy import copy
from pathlib import Path
from typing import Dict, Iterable, List, Tuple

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string, get_column_letter

SETTINGS_SHEET = "Settings"
PARAM_HEADER_ROW_1 = {
    "A1": "BasicField",
    "O1": "Alias",
    "P1": "ExtField",
}
PARAM_HEADER_ROW_2 = {
    "A2": "固定数据类型，不可修改",
    "O2": "本列必须包含，值可以为空",
    "P2": "char*",
    "Q2": "char*",
    "R2": "char*",
}
PARAM_HEADER_ROW_3 = {
    "A3": "参数序号",
    "B3": "隐藏",
    "C3": "只读",
    "D3": "存储",
    "E3": "复位",
    "F3": "运行时写入",
    "G3": "限制",
    "H3": "小数点",
    "I3": "符号",
    "J3": "浮点",
    "K3": "备用属性",
    "L3": "最大值",
    "M3": "最小值",
    "N3": "默认值",
    "O3": "参数别名",
    "P3": "Name",
    "Q3": "Unit",
    "R3": "Desc",
    "S3": "计算默认值",
}
PARAM_ID_RE = re.compile(r"^(?P<prefix>\d{3})-(?P<suffix>\d{3})$")
FRIENDLY_KEYS = {
    "id": "A",
    "parameter_id": "A",
    "hidden": "B",
    "readonly": "C",
    "read_only": "C",
    "storage": "D",
    "store": "D",
    "reset": "E",
    "runtime_write": "F",
    "limit": "G",
    "decimals": "H",
    "decimal": "H",
    "signed": "I",
    "float": "J",
    "reserved": "K",
    "max": "L",
    "min": "M",
    "default": "N",
    "alias": "O",
    "name": "P",
    "unit": "Q",
    "desc": "R",
    "description": "R",
    "calculated_default": "S",
    "computed_default": "S",
}
PREFERRED_COLUMN_NAMES = {
    "A": "id",
    "B": "hidden",
    "C": "read_only",
    "D": "storage",
    "E": "reset",
    "F": "runtime_write",
    "G": "limit",
    "H": "decimals",
    "I": "signed",
    "J": "float",
    "K": "reserved",
    "L": "max",
    "M": "min",
    "N": "default",
    "O": "alias",
    "P": "name",
    "Q": "unit",
    "R": "desc",
    "S": "calculated_default",
}
DEFAULT_PARAMETER_VALUES = {
    "B": 0,
    "C": 0,
    "D": 1,
    "E": 0,
    "F": 0,
    "G": 1,
    "H": 0,
    "I": 0,
    "J": 0,
    "K": 0,
    "L": 65535,
    "M": 0,
    "N": 100,
}
REQUIRED_PARAMETER_COLUMNS = ["O"]
DEFAULT_LIST_COLUMNS = ["A", "O", "N"]

warnings.filterwarnings(
    "ignore",
    message="Cannot parse header or footer so it will be ignored",
)


class ValidationError(Exception):
    pass


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("workbook", help="Path to ParameterTable.xlsx")
    parser.add_argument(
        "--output",
        help="Output workbook path. Defaults to overwriting the input workbook.",
    )

    subparsers = parser.add_subparsers(dest="command", required=True)

    add_macro = subparsers.add_parser("add-macro", help="Add a macro definition")
    add_macro.add_argument("--name", required=True)
    add_macro.add_argument("--value", required=True)
    add_macro.add_argument("--description", default="")

    remove_macro = subparsers.add_parser("remove-macro", help="Remove a macro definition")
    remove_macro.add_argument("--name", required=True)

    add_page = subparsers.add_parser("add-page", help="Add a parameter page")
    add_page.add_argument("--page-name", required=True)
    add_page.add_argument("--page-prefix", type=int, required=True)
    add_page.add_argument("--template-sheet", default="BASE")
    add_page.add_argument(
        "--copy-parameters",
        action="store_true",
        help="Keep parameter rows from the template sheet instead of creating an empty page.",
    )

    remove_page = subparsers.add_parser("remove-page", help="Remove a parameter page")
    remove_page.add_argument("--page-name", required=True)

    add_param = subparsers.add_parser("add-parameter", help="Add a parameter row")
    add_param.add_argument("--sheet", required=True)
    add_param.add_argument(
        "--suffix",
        default="auto",
        help="YYY in XXX-YYY. Use an integer or 'auto'.",
    )
    add_param.add_argument(
        "--prefix",
        type=int,
        help="XXX in XXX-YYY. Required when the page has no existing parameters.",
    )
    add_param.add_argument(
        "--json",
        help="Inline JSON object or @path/to/json file with row values.",
    )
    add_param.add_argument(
        "--set",
        action="append",
        default=[],
        metavar="KEY=VALUE",
        help="Set a row field. Keys can be Excel columns or friendly names like alias/default/unit.",
    )

    remove_param = subparsers.add_parser("remove-parameter", help="Remove a parameter row")
    remove_param.add_argument("--sheet", required=True)
    remove_param.add_argument("--id")
    remove_param.add_argument("--alias")

    edit_param = subparsers.add_parser("edit-parameter", help="Edit an existing parameter row")
    edit_param.add_argument("--sheet", required=True)
    edit_param.add_argument("--id")
    edit_param.add_argument("--alias")
    edit_param.add_argument(
        "--json",
        help="Inline JSON object or @path/to/json file with row values.",
    )
    edit_param.add_argument(
        "--set",
        action="append",
        default=[],
        metavar="KEY=VALUE",
        help="Set a row field. Keys can be Excel columns or friendly names like alias/default/unit.",
    )

    list_pages_parser = subparsers.add_parser("list-pages", help="List parameter page names")
    list_pages_parser.add_argument(
        "--json",
        action="store_true",
        help="Print the result as JSON.",
    )

    list_params = subparsers.add_parser("list-parameters", help="List parameters in a page")
    list_params.add_argument("--sheet", required=True)
    list_params.add_argument(
        "--columns",
        action="append",
        default=[],
        metavar="COL1,COL2",
        help="Columns to display. Use Excel columns or friendly names like id, alias, default.",
    )
    list_params.add_argument(
        "--json",
        action="store_true",
        help="Print the result as JSON.",
    )

    get_param = subparsers.add_parser("get-parameter", help="Show all properties for one parameter")
    get_param.add_argument("--sheet", required=True)
    get_param.add_argument("--id")
    get_param.add_argument("--alias")
    get_param.add_argument(
        "--json",
        action="store_true",
        help="Print the result as JSON.",
    )

    subparsers.add_parser("validate", help="Validate workbook structure and numbering")
    return parser.parse_args()


def normalize_output_path(workbook: Path, output: str | None) -> Path:
    return Path(output) if output else workbook


def load_wb(path: Path):
    return load_workbook(path)


def save_wb(wb, path: Path) -> None:
    wb.save(path)


def parameter_sheets(wb) -> List:
    return [ws for ws in wb.worksheets if ws.title != SETTINGS_SHEET]


def get_sheet_or_raise(wb, name: str):
    if name not in wb.sheetnames:
        raise ValidationError(f"Sheet '{name}' not found.")
    return wb[name]


def find_page_end_row(ws) -> int:
    matches = [cell.row for cell in ws["A"] if cell.value == "Page End"]
    if not matches:
        raise ValidationError(f"Sheet '{ws.title}' is missing 'Page End' in column A.")
    if len(matches) > 1:
        raise ValidationError(f"Sheet '{ws.title}' has multiple 'Page End' rows.")
    return matches[0]


def iter_parameter_rows(ws) -> Iterable[int]:
    end_row = find_page_end_row(ws)
    for row in range(4, end_row):
        yield row


def parse_param_id(value: str) -> Tuple[int, int]:
    match = PARAM_ID_RE.match(str(value).strip())
    if not match:
        raise ValidationError(f"Invalid parameter id '{value}'. Expected XXX-YYY.")
    return int(match.group("prefix")), int(match.group("suffix"))


def format_param_id(prefix: int, suffix: int) -> str:
    return f"{prefix:03d}-{suffix:03d}"


def ensure_prefix_unused(wb, prefix: int, ignore_sheet: str | None = None) -> None:
    for ws in parameter_sheets(wb):
        if ignore_sheet and ws.title == ignore_sheet:
            continue
        for row in iter_parameter_rows(ws):
            value = ws[f"A{row}"].value
            if not value:
                continue
            current_prefix, _ = parse_param_id(str(value))
            if current_prefix == prefix:
                raise ValidationError(
                    f"Prefix {prefix:03d} is already used by sheet '{ws.title}'."
                )
            break


def find_macro_row(ws, name: str) -> int | None:
    row = 2
    while True:
        value = ws[f"F{row}"].value
        if value == name:
            return row
        if value in (None, ""):
            return None
        row += 1


def find_next_macro_row(ws) -> int:
    row = 2
    while ws[f"F{row}"].value not in (None, ""):
        row += 1
    return row


def copy_row_style(ws, source_row: int, target_row: int, max_col: int | None = None) -> None:
    max_col = max_col or ws.max_column
    ws.row_dimensions[target_row].height = ws.row_dimensions[source_row].height
    for col in range(1, max_col + 1):
        source = ws.cell(source_row, col)
        target = ws.cell(target_row, col)
        if source.has_style:
            target._style = copy(source._style)
        if source.number_format:
            target.number_format = copy(source.number_format)
        if source.font:
            target.font = copy(source.font)
        if source.fill:
            target.fill = copy(source.fill)
        if source.border:
            target.border = copy(source.border)
        if source.alignment:
            target.alignment = copy(source.alignment)
        if source.protection:
            target.protection = copy(source.protection)
    for key, dimension in ws.column_dimensions.items():
        if dimension.width is not None:
            ws.column_dimensions[key].width = dimension.width


def clear_row_values(ws, row: int, max_col: int | None = None) -> None:
    max_col = max_col or ws.max_column
    for col in range(1, max_col + 1):
        ws.cell(row, col).value = None


def determine_sheet_prefix(ws) -> int | None:
    for row in iter_parameter_rows(ws):
        value = ws[f"A{row}"].value
        if value not in (None, ""):
            prefix, _ = parse_param_id(str(value))
            return prefix
    return None


def existing_suffixes(ws) -> List[int]:
    suffixes = []
    for row in iter_parameter_rows(ws):
        value = ws[f"A{row}"].value
        if value in (None, ""):
            continue
        _, suffix = parse_param_id(str(value))
        suffixes.append(suffix)
    return suffixes


def resolve_row_payload(json_arg: str | None, sets: List[str]) -> Dict[str, str]:
    payload: Dict[str, str] = {}
    if json_arg:
        if json_arg.startswith("@"):
            payload = json.loads(Path(json_arg[1:]).read_text(encoding="utf-8"))
        else:
            payload = json.loads(json_arg)
        if not isinstance(payload, dict):
            raise ValidationError("Row JSON must be an object.")
    for item in sets:
        if "=" not in item:
            raise ValidationError(f"Invalid --set '{item}'. Expected KEY=VALUE.")
        key, value = item.split("=", 1)
        payload[key] = value
    normalized: Dict[str, str] = {}
    for key, value in payload.items():
        if not isinstance(key, str):
            raise ValidationError("Row JSON keys must be strings.")
        key_upper = key.upper()
        column = FRIENDLY_KEYS.get(key.lower())
        if column is None and re.fullmatch(r"[A-Z]{1,3}", key_upper):
            column = key_upper
        if column is None:
            raise ValidationError(f"Unsupported field '{key}'.")
        normalized[column] = value
    return normalized


def add_macro(wb, name: str, value: str, description: str) -> None:
    ws = get_sheet_or_raise(wb, SETTINGS_SHEET)
    if find_macro_row(ws, name) is not None:
        raise ValidationError(f"Macro '{name}' already exists.")
    row = find_next_macro_row(ws)
    donor = max(2, row - 1)
    if row > 2:
        copy_row_style(ws, donor, row, max_col=8)
        clear_row_values(ws, row, max_col=8)
    ws[f"F{row}"] = name
    ws[f"G{row}"] = value
    ws[f"H{row}"] = description


def remove_macro(wb, name: str) -> None:
    ws = get_sheet_or_raise(wb, SETTINGS_SHEET)
    row = find_macro_row(ws, name)
    if row is None:
        raise ValidationError(f"Macro '{name}' not found.")
    ws.delete_rows(row, 1)


def add_page(wb, page_name: str, page_prefix: int, template_sheet: str, copy_parameters: bool) -> None:
    if page_name in wb.sheetnames:
        raise ValidationError(f"Sheet '{page_name}' already exists.")
    if page_name == SETTINGS_SHEET:
        raise ValidationError("Parameter page name cannot be 'Settings'.")
    ensure_prefix_unused(wb, page_prefix)
    source = get_sheet_or_raise(wb, template_sheet)
    if source.title == SETTINGS_SHEET:
        raise ValidationError("Cannot use Settings as a parameter page template.")

    cloned = wb.copy_worksheet(source)
    cloned.title = page_name

    if not copy_parameters:
        end_row = find_page_end_row(cloned)
        param_count = max(0, end_row - 4)
        if param_count:
            cloned.delete_rows(4, param_count)
    else:
        for row in iter_parameter_rows(cloned):
            current = cloned[f"A{row}"].value
            if current in (None, ""):
                continue
            _, suffix = parse_param_id(str(current))
            cloned[f"A{row}"] = format_param_id(page_prefix, suffix)


def remove_page(wb, page_name: str) -> None:
    if page_name == SETTINGS_SHEET:
        raise ValidationError("Settings sheet cannot be removed.")
    ws = get_sheet_or_raise(wb, page_name)
    wb.remove(ws)


def pick_style_donor_row(wb, ws) -> Tuple[int, int]:
    end_row = find_page_end_row(ws)
    if end_row > 4:
        return end_row - 1, ws.max_column
    for candidate in parameter_sheets(wb):
        if candidate.title == ws.title:
            continue
        candidate_end = find_page_end_row(candidate)
        if candidate_end > 4:
            return 4, candidate.max_column
    return 4, max(ws.max_column, 19)


def add_parameter(wb, sheet: str, suffix_raw: str, prefix_override: int | None, payload: Dict[str, str]) -> str:
    ws = get_sheet_or_raise(wb, sheet)
    if ws.title == SETTINGS_SHEET:
        raise ValidationError("Parameters cannot be added to Settings.")

    prefix = determine_sheet_prefix(ws)
    if prefix is None:
        if prefix_override is None:
            raise ValidationError("Empty parameter pages require --prefix.")
        ensure_prefix_unused(wb, prefix_override, ignore_sheet=ws.title)
        prefix = prefix_override
    elif prefix_override is not None and prefix_override != prefix:
        raise ValidationError(
            f"Sheet '{sheet}' already uses prefix {prefix:03d}; received {prefix_override:03d}."
        )

    suffixes = existing_suffixes(ws)
    if suffix_raw == "auto":
        suffix = 0 if not suffixes else max(suffixes) + 1
    else:
        suffix = int(suffix_raw)
    if suffix > 127 or suffix < 0:
        raise ValidationError("Suffix must be between 0 and 127.")
    if suffix in suffixes:
        raise ValidationError(f"Suffix {suffix:03d} already exists in sheet '{sheet}'.")

    end_row = find_page_end_row(ws)
    ws.insert_rows(end_row, 1)
    donor_row, donor_max_col = pick_style_donor_row(wb, ws)
    copy_row_style(ws, donor_row, end_row, max_col=donor_max_col)
    clear_row_values(ws, end_row, max_col=max(donor_max_col, ws.max_column))

    param_id = format_param_id(prefix, suffix)
    ws[f"A{end_row}"] = param_id
    for column, value in DEFAULT_PARAMETER_VALUES.items():
        ws[f"{column}{end_row}"] = value
    for column, value in payload.items():
        if column == "A":
            continue
        ws[f"{column}{end_row}"] = value
    for column in REQUIRED_PARAMETER_COLUMNS:
        if ws[f"{column}{end_row}"].value in (None, ""):
            raise ValidationError(
                f"New parameter requires column {column} to be set."
            )
    return param_id


def find_parameter_row(ws, param_id: str | None, alias: str | None) -> int:
    if bool(param_id) == bool(alias):
        raise ValidationError("Provide exactly one of --id or --alias.")
    if ws.title == SETTINGS_SHEET:
        raise ValidationError("Parameters cannot be edited on Settings.")

    for row in iter_parameter_rows(ws):
        row_id = ws[f"A{row}"].value
        row_alias = ws[f"O{row}"].value
        if param_id and str(row_id) == param_id:
            return row
        if alias and str(row_alias) == alias:
            return row

    target = param_id if param_id else alias
    raise ValidationError(f"Parameter '{target}' not found in sheet '{ws.title}'.")


def edit_parameter(
    wb,
    sheet: str,
    param_id: str | None,
    alias: str | None,
    payload: Dict[str, str],
) -> str:
    if not payload:
        raise ValidationError("Provide at least one field with --json or --set.")

    ws = get_sheet_or_raise(wb, sheet)
    row = find_parameter_row(ws, param_id, alias)
    current_id = ws[f"A{row}"].value

    for column, value in payload.items():
        if column == "A":
            raise ValidationError("Editing column A / parameter id is not supported.")
        ws[f"{column}{row}"] = value

    for column in REQUIRED_PARAMETER_COLUMNS:
        if ws[f"{column}{row}"].value in (None, ""):
            raise ValidationError(
                f"Edited parameter requires column {column} to be set."
            )

    return str(current_id)


def remove_parameter(wb, sheet: str, param_id: str | None, alias: str | None) -> None:
    ws = get_sheet_or_raise(wb, sheet)
    row = find_parameter_row(ws, param_id, alias)
    ws.delete_rows(row, 1)


def list_pages(wb) -> List[str]:
    return [ws.title for ws in parameter_sheets(wb)]


def resolve_query_columns(column_args: List[str]) -> List[str]:
    if not column_args:
        return list(DEFAULT_LIST_COLUMNS)

    resolved: List[str] = []
    seen = set()
    for item in column_args:
        for raw_key in item.split(","):
            key = raw_key.strip()
            if not key:
                continue
            key_upper = key.upper()
            column = FRIENDLY_KEYS.get(key.lower())
            if column is None and re.fullmatch(r"[A-Z]{1,3}", key_upper):
                column = key_upper
            if column is None:
                raise ValidationError(f"Unsupported query column '{key}'.")
            if column not in seen:
                resolved.append(column)
                seen.add(column)
    if not resolved:
        raise ValidationError("Provide at least one query column.")
    return resolved


def column_label(column: str) -> str:
    return PREFERRED_COLUMN_NAMES.get(column, column.lower())


def cell_to_text(value) -> str:
    return "" if value is None else str(value)


def parameter_row_to_dict(ws, row: int, columns: List[str]) -> Dict[str, str]:
    return {
        column_label(column): cell_to_text(ws[f"{column}{row}"].value)
        for column in columns
    }


def parameter_detail_columns(ws) -> List[str]:
    return [get_column_letter(col) for col in range(1, ws.max_column + 1)]


def list_parameters(wb, sheet: str, columns: List[str]) -> List[Dict[str, str]]:
    ws = get_sheet_or_raise(wb, sheet)
    if ws.title == SETTINGS_SHEET:
        raise ValidationError("Settings is not a parameter page.")

    parameters: List[Dict[str, str]] = []
    for row in iter_parameter_rows(ws):
        param_id = ws[f"A{row}"].value
        if param_id in (None, "", "Page End"):
            continue
        parameters.append(parameter_row_to_dict(ws, row, columns))
    return parameters


def get_parameter(wb, sheet: str, param_id: str | None, alias: str | None) -> Dict[str, str]:
    ws = get_sheet_or_raise(wb, sheet)
    row = find_parameter_row(ws, param_id, alias)
    return parameter_row_to_dict(ws, row, parameter_detail_columns(ws))


def print_pages(pages: List[str]) -> None:
    if not pages:
        print("No parameter pages found.")
        return
    for page in pages:
        print(page)


def print_parameters(parameters: List[Dict[str, str]]) -> None:
    if not parameters:
        print("No parameters found.")
        return
    for item in parameters:
        print("\t".join(f"{key}={value}" for key, value in item.items()))


def print_parameter_detail(parameter: Dict[str, str]) -> None:
    for key, value in parameter.items():
        print(f"{key}={value}")


def print_json(data) -> None:
    print(json.dumps(data, ensure_ascii=False, indent=2))


def validate_workbook(wb) -> Tuple[List[str], List[str]]:
    errors: List[str] = []
    warnings: List[str] = []
    seen_prefixes: Dict[int, str] = {}

    if SETTINGS_SHEET not in wb.sheetnames:
        errors.append("Missing Settings sheet.")
        return errors, warnings

    for ws in parameter_sheets(wb):
        for cell_ref, expected in PARAM_HEADER_ROW_1.items():
            if ws[cell_ref].value != expected:
                errors.append(f"{ws.title}: {cell_ref} must be '{expected}'.")
        for cell_ref, expected in PARAM_HEADER_ROW_2.items():
            if ws[cell_ref].value != expected:
                errors.append(f"{ws.title}: {cell_ref} must be '{expected}'.")
        for cell_ref, expected in PARAM_HEADER_ROW_3.items():
            if ws[cell_ref].value != expected:
                errors.append(f"{ws.title}: {cell_ref} must be '{expected}'.")

        try:
            end_row = find_page_end_row(ws)
        except ValidationError as exc:
            errors.append(str(exc))
            continue

        if end_row < 4:
            errors.append(f"{ws.title}: 'Page End' must be after row 3.")
            continue

        page_prefix = None
        suffixes: List[int] = []

        for row in range(4, end_row):
            value = ws[f"A{row}"].value
            if value in (None, ""):
                errors.append(f"{ws.title}: row {row} is missing parameter id in column A.")
                continue
            try:
                prefix, suffix = parse_param_id(str(value))
            except ValidationError as exc:
                errors.append(f"{ws.title}: row {row}: {exc}")
                continue

            if page_prefix is None:
                page_prefix = prefix
            elif prefix != page_prefix:
                errors.append(
                    f"{ws.title}: row {row} uses prefix {prefix:03d}, expected {page_prefix:03d}."
                )

            if suffix > 127:
                errors.append(f"{ws.title}: row {row} suffix {suffix:03d} exceeds 127.")
            if suffix in suffixes:
                errors.append(f"{ws.title}: duplicate suffix {suffix:03d}.")
            suffixes.append(suffix)

            if ws[f"O{row}"].value in (None, ""):
                errors.append(f"{ws.title}: row {row} column O (Alias) must not be empty.")

        if page_prefix is None:
            warnings.append(f"{ws.title}: page has no parameter rows.")
            continue

        owner = seen_prefixes.get(page_prefix)
        if owner and owner != ws.title:
            errors.append(
                f"Sheet '{ws.title}' reuses prefix {page_prefix:03d}, already used by '{owner}'."
            )
        else:
            seen_prefixes[page_prefix] = ws.title

        ordered = sorted(set(suffixes))
        expected = list(range(ordered[0], ordered[-1] + 1)) if ordered else []
        if ordered and ordered != expected:
            missing = [f"{item:03d}" for item in expected if item not in ordered]
            warnings.append(
                f"{ws.title}: parameter suffixes are not continuous; missing {', '.join(missing)}."
            )
        if ordered and ordered[0] != 0:
            warnings.append(f"{ws.title}: parameter suffixes do not start at 000.")

    return errors, warnings


def print_validation(errors: List[str], warnings: List[str]) -> None:
    if errors:
        print("Validation errors:")
        for item in errors:
            print(f"- {item}")
    if warnings:
        print("Validation warnings:")
        for item in warnings:
            print(f"- {item}")
    if not errors and not warnings:
        print("Validation passed with no issues.")


def main() -> int:
    args = parse_args()
    workbook_path = Path(args.workbook)
    output_path = normalize_output_path(workbook_path, args.output)

    try:
        wb = load_wb(workbook_path)

        if args.command == "add-macro":
            add_macro(wb, args.name, args.value, args.description)
            save_wb(wb, output_path)
            print(f"Added macro '{args.name}'.")
        elif args.command == "remove-macro":
            remove_macro(wb, args.name)
            save_wb(wb, output_path)
            print(f"Removed macro '{args.name}'.")
        elif args.command == "add-page":
            add_page(wb, args.page_name, args.page_prefix, args.template_sheet, args.copy_parameters)
            save_wb(wb, output_path)
            print(f"Added page '{args.page_name}' with prefix {args.page_prefix:03d}.")
        elif args.command == "remove-page":
            remove_page(wb, args.page_name)
            save_wb(wb, output_path)
            print(f"Removed page '{args.page_name}'.")
        elif args.command == "add-parameter":
            payload = resolve_row_payload(args.json, args.set)
            param_id = add_parameter(wb, args.sheet, args.suffix, args.prefix, payload)
            save_wb(wb, output_path)
            print(f"Added parameter '{param_id}' to sheet '{args.sheet}'.")
        elif args.command == "remove-parameter":
            remove_parameter(wb, args.sheet, args.id, args.alias)
            save_wb(wb, output_path)
            print(f"Removed parameter from sheet '{args.sheet}'.")
        elif args.command == "edit-parameter":
            payload = resolve_row_payload(args.json, args.set)
            param_id = edit_parameter(wb, args.sheet, args.id, args.alias, payload)
            save_wb(wb, output_path)
            print(f"Edited parameter '{param_id}' in sheet '{args.sheet}'.")
        elif args.command == "list-pages":
            pages = list_pages(wb)
            if args.json:
                print_json(pages)
            else:
                print_pages(pages)
        elif args.command == "list-parameters":
            columns = resolve_query_columns(args.columns)
            parameters = list_parameters(wb, args.sheet, columns)
            if args.json:
                print_json(parameters)
            else:
                print_parameters(parameters)
        elif args.command == "get-parameter":
            parameter = get_parameter(wb, args.sheet, args.id, args.alias)
            if args.json:
                print_json(parameter)
            else:
                print_parameter_detail(parameter)
        elif args.command == "validate":
            errors, warnings = validate_workbook(wb)
            print_validation(errors, warnings)
            return 1 if errors else 0
        else:
            raise ValidationError(f"Unsupported command '{args.command}'.")
        return 0
    except ValidationError as exc:
        print(f"Error: {exc}", file=sys.stderr)
        return 2


if __name__ == "__main__":
    sys.exit(main())
